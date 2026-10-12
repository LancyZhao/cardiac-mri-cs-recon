"""Format the provided-paper translation as a separate, editable attachment."""
import argparse
import json
import re
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from build_thesis_docx import configure, add_field


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    args = ap.parse_args()
    text = args.source.read_text()
    body = text.split("## 中文译文", 1)[1].split("## 译文核对说明", 1)[0]
    count = len(re.findall(r"[\u4e00-\u9fff]", body))
    if not 5000 <= count <= 6000:
        raise SystemExit(f"translation body must be 5000-6000 Chinese characters, got {count}")
    pypandoc.convert_text(text, "docx", format="markdown+tex_math_dollars",
                         outputfile=str(args.output), extra_args=["--standalone"])
    doc = Document(args.output)
    configure(doc)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(2.75), Cm(2.54)
        section.left_margin, section.right_margin = Cm(3.57), Cm(2.77)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent = Pt(0)
        add_field(footer, " PAGE ")
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            p.paragraph_format.page_break_before = False
            p.paragraph_format.keep_with_next = True
        if p._p.xpath(".//m:oMathPara"):
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(args.output)
    print(json.dumps({"translation_chinese_chars": count,
                      "paragraphs": len(doc.paragraphs),
                      "math_nodes": len(doc.element.body.xpath(".//m:oMath"))}))


if __name__ == "__main__":
    main()
