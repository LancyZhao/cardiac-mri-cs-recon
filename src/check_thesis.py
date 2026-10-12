"""Static checks for the thesis draft and generated Word document."""
from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from fill_thesis_from_report import measured_tables, replace_tables


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--draft", type=Path, required=True)
    ap.add_argument("--docx", type=Path, required=True)
    ap.add_argument("--results", type=Path)
    ap.add_argument("--output", type=Path)
    args = ap.parse_args()
    text = args.draft.read_text()
    body = text[text.index("## 第一章"):text.index("## 参考文献")]
    doc = Document(args.docx)
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    table_captions = [
        p.text for p in doc.paragraphs
        if re.match(r"^表\d+-\d+", p.text)
    ]
    image_count = len(doc.inline_shapes)
    math_count = 0
    with zipfile.ZipFile(args.docx) as archive:
        math_count = archive.read("word/document.xml").count(b"<m:oMath")
    checks = {
        "body_chinese_chars": sum("\u4e00" <= c <= "\u9fff" for c in body),
        "references": len(re.findall(r"^\[\d+\]", text, re.M)),
        "placeholders": len(re.findall(r"待填写|X dB|实验完成后填写|预期：", text)),
        "old_lr_mentions": text.count("2\\times 10^{-4}") + text.count("lr=2e-4"),
        "old_param_mentions": text.count("0.01M"),
        "docx_bytes": args.docx.stat().st_size,
        "docx_paragraphs": len(doc.paragraphs),
        "docx_tables": len(doc.tables),
        "docx_images": image_count,
        "docx_math_nodes": math_count,
        "table_captions": table_captions,
        "heading_count": sum(p.style.name.startswith("Heading") for p in doc.paragraphs),
    }
    if args.results:
        p = args.results / "thesis_results.json"
        checks["measured_report"] = p.exists()
        if p.exists():
            report = json.loads(p.read_text())
            checks["report_status"] = report.get("status")
            checks["tables_match_measured_json"] = (
                replace_tables(text, measured_tables(report)) == text)
            verification = json.loads((args.results / "reevaluation.json").read_text())
            checks["checkpoint_reevaluation_passed"] = verification.get("passed")
    checks["idempotent_table_fill"] = checks.get("tables_match_measured_json", False)
    checks["sections"] = len(doc.sections)
    checks["margins_match_template"] = all(
        abs(getattr(s, field) - Cm(value)) < 1000
        for s in doc.sections
        for field, value in (("top_margin", 2.75), ("bottom_margin", 2.54),
                             ("left_margin", 3.57), ("right_margin", 2.77)))
    checks["body_style_12pt"] = doc.styles["Normal"].font.size == Pt(12)
    checks["body_line_spacing_20pt"] = doc.styles["Normal"].paragraph_format.line_spacing == Pt(20)
    checks["broken_headings"] = bool(re.search(r"[^\n#]#{3,4}\s+\d", text))
    refs = set(map(int, re.findall(r"^\[(\d+)\]", text, re.M)))
    cites = set()
    for match in re.findall(r"\[([1-9]\d*(?:[,-][1-9]\d*)*)\]", body):
        for token in match.split(","):
            if "-" in token:
                lo, hi = map(int, token.split("-"))
                cites.update(range(lo, hi + 1))
            else:
                cites.add(int(token))
    checks["uncited_references"] = sorted(refs - cites)
    checks["undefined_citations"] = sorted(cites - refs)
    checks["docx_matches_body"] = all(
        re.sub(r"\s+", "", sentence) in re.sub(r"\s+", "", doc_text)
        for sentence in ("本文围绕心脏MRI加速重建问题，完成了以下研究工作：",
                         "感谢指导教师孙学斌老师，感谢学院为毕业设计提供的学习条件与组织支持。"))
    actual_tables = [
        [[re.sub(r"\s+", "", c.text) for c in row.cells] for row in t.rows]
        for t in doc.tables]
    checks["word_tables_match_json"] = True
    if args.results:
        for table in measured_tables(report).values():
            expected = [
                [re.sub(r"\s+", "", c) for c in line.strip("|").split("|")]
                for line in table.splitlines() if line.startswith("|")
                and not re.fullmatch(r"[|:\-\s]+", line)]
            if expected not in actual_tables:
                checks["word_tables_match_json"] = False
    equations = re.findall(r"\((\d+)-(\d+)\)\s*$", doc_text, re.M)
    checks["equation_number_count"] = len(equations)
    checks["equation_chapters_valid"] = bool(equations) and all(
        1 <= int(chapter) <= 6 for chapter, number in equations)
    failures = []
    if checks["body_chinese_chars"] < 15000:
        failures.append("正文汉字不足15000")
    if checks["references"] < 15:
        failures.append("参考文献少于15篇")
    if checks["old_lr_mentions"] or checks["old_param_mentions"]:
        failures.append("仍有旧训练参数或参数量")
    if checks["docx_bytes"] < 50000:
        failures.append("Word文件过小，可能未生成完整正文")
    if checks["docx_images"] < 5:
        failures.append("论文未嵌入全部研究图")
    if checks["docx_math_nodes"] < 30:
        failures.append("公式数量异常偏少")
    if len(checks["table_captions"]) < 8:
        failures.append("表题数量不足")
    if "表5-4  全图口径PSNR对比（dB，按切片平均）" not in text:
        failures.append("主结果表题缺失")
    for key in ("tables_match_measured_json", "checkpoint_reevaluation_passed",
                "margins_match_template", "body_style_12pt",
                "body_line_spacing_20pt", "docx_matches_body", "word_tables_match_json",
                "equation_chapters_valid"):
        if not checks.get(key):
            failures.append(key)
    for key in ("placeholders", "broken_headings", "uncited_references", "undefined_citations"):
        if checks[key]:
            failures.append(key)
    checks["passed"] = not failures
    checks["failures"] = failures
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    if args.output:
        args.output.write_text(json.dumps(checks, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit("; ".join(failures))


if __name__ == "__main__":
    main()
