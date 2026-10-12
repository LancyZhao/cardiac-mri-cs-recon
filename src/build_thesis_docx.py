"""Convert Markdown math with Pandoc, then apply the supplied university rules."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import argparse
import json
import re
import shutil
import tempfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.run import Run


def font_style(style, size=12, chinese="宋体", bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rp = style.element.get_or_add_rPr()
    fonts = rp.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rp.append(fonts)
    for k,v in (("eastAsia", chinese), ("ascii","Times New Roman"),
                ("hAnsi","Times New Roman")):
        fonts.set(qn("w:"+k), v)


def configure(doc):
    for style in doc.styles:
        if style.type == 1:
            font_style(style)
    for name in ("Normal", "Body Text", "First Paragraph"):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        p = style.paragraph_format
        p.line_spacing = Pt(20)
        p.space_before = p.space_after = Pt(0)
        p.first_line_indent = Pt(24)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level,size,spacing in ((1,15,30),(2,14,18),(3,14,12)):
        s = doc.styles[f"Heading {level}"]
        font_style(s, size, "黑体", True)
        f = s.paragraph_format
        f.first_line_indent = Pt(0)
        f.space_before = f.space_after = Pt(spacing)
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        f.keep_with_next = True
        f.page_break_before = level == 1
    for name in ("Caption", "Image Caption", "Table Caption"):
        if name in doc.styles:
            font_style(doc.styles[name], 10.5)
    for level in range(1, 4):
        name = f"TOC {level}"
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        font_style(doc.styles[name], 12)
        f = doc.styles[name].paragraph_format
        f.line_spacing = Pt(20)
        f.space_before = f.space_after = Pt(0)
        f.first_line_indent = Pt(0)
        f.left_indent = Pt(24 * (level - 1))
        f.tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def add_field(p, instruction, result=""):
    for kind in ("begin", "instruction", "separate", "result", "end"):
        r = OxmlElement("w:r")
        if kind == "instruction":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instruction
        elif kind == "result":
            el = OxmlElement("w:t")
            el.text = result
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r.append(el)
        p._p.append(r)


def section_before(p):
    sect = deepcopy(Document().sections[0]._sectPr)
    q = p.insert_paragraph_before()
    q._p.get_or_add_pPr().append(sect)


def cover_elements(title, month, template=None):
    if template:
        cover = Document(template)
        abstract = next(
            (p for p in cover.paragraphs if p.text.strip().startswith("摘  要")),
            None,
        )
        if abstract is not None:
            started = False
            for child in list(cover.element.body):
                if child is abstract._p:
                    started = True
                if started and child.tag != qn("w:sectPr"):
                    cover.element.body.remove(child)
        values = {"学    院": "电气自动化与信息工程学院", "专    业": "物联网工程",
                  "年    级": "2023级", "姓    名": "赵兰熙",
                  "学    号": "3023234248", "指导教师": "孙学斌"}
        for p in cover.paragraphs:
            if p.text.strip().startswith("题目："):
                p.text = "题目：" + title.replace("心脏", "\n心脏", 1)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.left_indent = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(18)
            for label, value in values.items():
                if p.text.strip() == label:
                    p.text = label + "  " + value
            if p.text.strip() == "独创性声明":
                p.paragraph_format.page_break_before = True
                p.paragraph_format.keep_with_next = True
        return cover

    cover = Document()
    configure(cover)
    p = cover.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(70)
    r = p.add_run("本科生毕业论文" if not template else "天津大学本科生毕业论文")
    r.bold, r.font.size = True, Pt(24)
    p = cover.add_paragraph(title)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(50)
    for label,value in (
        ("学    院","电气自动化与信息工程学院"), ("专    业","物联网工程"),
        ("年    级","2023级"), ("姓    名","赵兰熙"), ("学    号","3023234248"),
        ("指导教师","孙学斌"),
    ):
        p = cover.add_paragraph(label+"    "+value)
        p.paragraph_format.left_indent = Cm(1.6)
        p.paragraph_format.space_after = Pt(10)
    p = cover.add_paragraph(month)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.add_page_break()
    p = cover.add_paragraph("独创性声明")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(22)
    p.paragraph_format.space_after = Pt(30)
    cover.add_paragraph(
        "本人声明：所呈交的毕业设计（论文），是本人在指导教师指导下，"
        "进行研究工作所取得的成果。除文中已经注明引用的内容外，"
        "本毕业设计（论文）中不包含任何他人已经发表或撰写过的研究成果。"
        "对本毕业设计（论文）所涉及的研究工作做出贡献的其他个人和集体，"
        "均已在论文中作了明确的说明。本毕业设计（论文）原创性声明的法律责任由本人承担。"
    )
    cover.add_paragraph("论文作者签名：_______________")
    cover.add_paragraph("日期：______年______月______日")
    cover.add_paragraph()
    cover.add_paragraph("本人声明：本毕业设计（论文）是本人指导学生完成的研究成果，已经审阅过论文的全部内容。")
    cover.add_paragraph("论文指导教师签名：_______________")
    cover.add_paragraph("日期：______年______月______日")
    return cover


def table_format(table):
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    count = len(table.columns)
    usable = 14.66
    widths = [4.7] + [(usable - 4.7)/(count-1)]*(count-1) if count > 2 else [5.0, 9.66]
    for column,width in zip(table.columns, widths):
        column.width = Cm(width)
    for i,row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j,cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "bottom", "left", "right"):
                b = OxmlElement("w:"+edge)
                line = edge == "top" and i == 0 or edge == "bottom" and i in (0,len(table.rows)-1)
                b.set(qn("w:val"), "single" if line else "nil")
                if line:
                    b.set(qn("w:sz"), "8" if i == 0 else "12")
                borders.append(b)
            cell._tc.get_or_add_tcPr().append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = Pt(15)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.bold = i == 0


def cached_toc(chapter, doc, page_map):
    toc = chapter.insert_paragraph_before()
    add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "")
    end = toc._p[-1]
    toc._p.remove(end)
    entries = [p for p in doc.paragraphs
               if p.style.name in ("Heading 1", "Heading 2", "Heading 3")
               and p.text not in ("摘 要", "ABSTRACT")]
    for entry in entries:
        level = int(entry.style.name[-1])
        p = chapter.insert_paragraph_before()
        p.style = doc.styles[f"TOC {level}"]
        key = re.sub(r"\s+", "", entry.text)
        number = str(page_map.get(key, ""))
        p.add_run(entry.text + "\t" + number)
    p._p.append(end)


def mark_citations(p):
    pattern = r"(\[[1-9]\d*(?:[,-][1-9]\d*)*\])"
    for run in list(p.runs):
        pieces = re.split(pattern, run.text)
        if len(pieces) == 1:
            continue
        for text in pieces:
            if not text:
                continue
            el = deepcopy(run._r)
            part = Run(el, p)
            part.text = text
            if re.fullmatch(pattern, text):
                part.font.superscript = True
            run._r.addprevious(el)
        p._p.remove(run._r)


def build(source, output, month=None, template=None, toc_map=None):
    text = source.read_text()
    text = re.sub(r"^# .+\n", "", text, count=1)
    text = text.replace("## 摘要（中文，300-400字）","## 摘  要")
    text = re.sub(r"^\s*---\s*$", "", text, flags=re.M)
    with tempfile.TemporaryDirectory() as td:
        interim = Path(td)/"body.docx"
        pypandoc.convert_text(
            text, "docx", format="markdown+tex_math_dollars",
            outputfile=str(interim), extra_args=[
                "--shift-heading-level-by=-1", "--standalone",
                "--resource-path="+str(source.parent),
            ],
        )
        doc = Document(interim)
    configure(doc)
    first = doc.paragraphs[0]
    section_before(first)
    cover = cover_elements(
        "基于展开式迭代网络的心脏MRI加速重建研究",
        month or f"{date.today().year}年{date.today().month}月",
        template,
    )
    elements = []
    for el in cover.element.body:
        if el.tag == qn("w:sectPr"):
            continue
        el = deepcopy(el)
        for sect in el.xpath(".//w:sectPr"):
            sect.getparent().remove(sect)
        elements.append(el)
    for el in elements:
        for blip in el.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            part = cover.part.related_parts.get(rid)
            if part is not None and hasattr(part, "blob"):
                new_rid, _ = doc.part.get_or_add_image(BytesIO(part.blob))
                blip.set(qn("r:embed"), new_rid)
    for el in reversed(elements):
        doc.element.body.insert(0,el)
    chapter = next(p for p in doc.paragraphs if p.text.startswith("第一章"))
    toc_title = chapter.insert_paragraph_before("目  录")
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.page_break_before = True
    toc_title.runs[0].bold = True
    toc_title.runs[0].font.size = Pt(22)
    cached_toc(chapter, doc, toc_map or {})
    section_before(chapter)
    for i, section in enumerate(doc.sections):
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(2.75), Cm(2.54)
        section.left_margin, section.right_margin = Cm(3.57), Cm(2.77)
        section.header_distance, section.footer_distance = Cm(1.5), Cm(1.75)
        section.start_type = WD_SECTION_START.NEW_PAGE
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        for p in section.header.paragraphs + section.footer.paragraphs:
            p.clear()
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if i == 2:
            h = section.header.paragraphs[0]
            h.add_run("天津大学2026届本科生毕业设计（论文）").font.size = Pt(10.5)
        if i:
            nums = OxmlElement("w:pgNumType")
            nums.set(qn("w:fmt"), "upperRoman" if i == 1 else "decimal")
            nums.set(qn("w:start"), "1")
            section._sectPr.append(nums)
            add_field(section.footer.paragraphs[0], " PAGE ")
            for run in section.footer.paragraphs[0].runs:
                run.font.size = Pt(9)
    chapter_no, equation_no = 0, 0
    in_refs = False
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and re.match(r"^第[一二三四五六]章", p.text):
            chapter_no = "一二三四五六".index(p.text[1]) + 1
            equation_no = 0
        if p.text in ("摘 要", "摘  要", "ABSTRACT"):
            if p.text != "ABSTRACT":
                p.text = "摘  要"
            p.style = doc.styles["Title"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.page_break_before = p.text == "ABSTRACT"
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(22)
        if p.text.startswith(("关键词", "KEY WORDS")):
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(20)
        if re.match(r"^[表图]\d+-\d+\s",p.text):
            p.style = doc.styles["Caption"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = p.text.startswith("表")
        if p.style.name.startswith("Heading"):
            # Pandoc collapses runs of spaces, so restore the required heading gap.
            p.text = re.sub(r"^(第[一二三四五六]章|\d+(?:\.\d+)+) +", r"\1  ", p.text)
        if p._p.xpath(".//m:oMathPara"):
            equation_no += 1
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.tab_stops.add_tab_stop(Cm(7.1), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.RIGHT)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(6)
            if chapter_no:
                math = p._p.xpath(".//m:oMathPara")[0]
                # Inline OMML on a centered tab keeps the equation number at the right margin.
                pos = p._p.index(math)
                before = OxmlElement("w:r")
                before.append(OxmlElement("w:tab"))
                p._p.insert(pos, before)
                for node in list(math):
                    if node.tag == qn("m:oMath"):
                        p._p.insert(p._p.index(math), deepcopy(node))
                p._p.remove(math)
                p.add_run(f"\t({chapter_no}-{equation_no})")
        if p.text == "参考文献":
            in_refs = True
        elif p.text.startswith(("附录", "致谢")):
            in_refs = False
        elif in_refs and p.text:
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-24)
        elif not p.style.name.startswith(("Heading", "TOC")):
            mark_citations(p)
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.keep_with_next = True
    for shape in doc.inline_shapes:
        if shape.width > Cm(14.6):
            ratio = Cm(14.6) / shape.width
            shape.height = int(shape.height * ratio)
            shape.width = Cm(14.6)
    for table in doc.tables:
        table_format(table)
    setting = OxmlElement("w:updateFields")
    setting.set(qn("w:val"), "true")
    doc.settings.element.append(setting)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        backup = output.with_suffix(".before-rebuild.docx")
        if not backup.exists():
            shutil.copy2(output, backup)
    doc.save(output)
    audit = {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
             "math_nodes": len(doc.element.body.xpath(".//m:oMath")),
             "images": len(doc.inline_shapes), "sections": len(doc.sections)}
    print(json.dumps(audit,ensure_ascii=False))
    return audit


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source",type=Path,required=True)
    ap.add_argument("--output",type=Path,required=True)
    ap.add_argument("--template", type=Path)
    ap.add_argument("--toc-map", type=Path)
    args=ap.parse_args()
    build(args.source,args.output,template=args.template,
          toc_map=json.loads(args.toc_map.read_text()) if args.toc_map else None)


if __name__=="__main__":
    main()
